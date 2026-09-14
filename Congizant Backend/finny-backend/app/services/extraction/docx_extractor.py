import logging
import re
from pathlib import Path
from typing import Union, List, Any
import docx
from app.services.extraction.base_extractor import BaseExtractor
from app.schemas.financial import RawExtractionData, RawTable, DocumentSource

logger = logging.getLogger(__name__)


class DocxExtractor(BaseExtractor):
    """Extractor for Microsoft Word (.docx) documents."""

    def __init__(self):
        super().__init__(file_type="docx")

    def extract(self, file_path: Union[str, Path]) -> RawExtractionData:
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"DOCX file not found: {file_path}")

        extracted_tables: List[RawTable] = []
        paragraph_texts: List[str] = []
        source = DocumentSource(filename=path.name, file_type=self.file_type)

        try:
            doc = docx.Document(path)

            # 1. Collect all paragraph text
            for p in doc.paragraphs:
                txt = p.text.strip()
                if txt:
                    paragraph_texts.append(txt)

            # 2. Extract structured tables
            for tbl_idx, table in enumerate(doc.tables, start=1):
                raw_rows: List[List[str]] = []
                for row in table.rows:
                    cleaned_row = [cell.text.strip() for cell in row.cells]
                    if any(c != "" for c in cleaned_row):
                        raw_rows.append(cleaned_row)

                if not raw_rows:
                    continue

                clean_tbl = self._clean_table(raw_rows, tbl_idx)
                if clean_tbl and clean_tbl.rows:
                    extracted_tables.append(clean_tbl)

            # 3. Fallback: if no tables found, parse structured text lines from paragraphs
            full_text = "\n".join(paragraph_texts)
            if not extracted_tables and full_text:
                text_table = self._extract_tables_from_text(full_text)
                if text_table and text_table.rows:
                    extracted_tables.append(text_table)

        except Exception as err:
            logger.exception("Failed to parse DOCX file: %s", path)
            raise ValueError(f"Failed to extract DOCX data: {str(err)}")

        return RawExtractionData(
            source=source,
            tables=extracted_tables,
            metadata={
                "total_tables_found": len(extracted_tables),
                "raw_text": "\n".join(paragraph_texts),
                "paragraphs_count": len(paragraph_texts),
            }
        )

    def _clean_table(self, rows_data: List[List[str]], table_idx: int) -> RawTable:
        """Sanitize raw docx table rows and determine header row."""
        if not rows_data:
            return RawTable(headers=[], rows=[])

        if len(rows_data) == 1:
            headers = [f"Col_{i+1}" for i in range(len(rows_data[0]))]
            return RawTable(headers=headers, rows=rows_data, page_number=table_idx)

        headers = rows_data[0]
        if all(h == "" for h in headers):
            headers = [f"Col_{i+1}" for i in range(len(headers))]
            data_rows = rows_data
        else:
            headers = [h if h != "" else f"Col_{i+1}" for i, h in enumerate(headers)]
            data_rows = rows_data[1:]

        return RawTable(headers=headers, rows=data_rows, page_number=table_idx)

    def _extract_tables_from_text(self, text: str) -> RawTable:
        """Parse structured tabular financial statements from plaintext lines."""
        rows: List[List[str]] = []
        lines = text.splitlines()

        should_reverse_cols = False
        detected_headers: List[str] = []
        for l in lines[:10]:
            years = re.findall(r"\b(20\d\d|19\d\d)\b", l)
            if len(years) >= 2:
                if int(years[0]) > int(years[-1]):
                    should_reverse_cols = True
                    detected_headers = ["Particulars"] + list(reversed(years))
                else:
                    detected_headers = ["Particulars"] + years
                break

        line_pattern = re.compile(
            r"^([A-Za-z\s'&/.-]+?)\s*[:=]?\s+([₹$€£]?\s*\(?[\d,]+(?:\.\d+)?\s*(?:M|million|B|billion|K|thousand|Cr|L)?\)?(?:\s+[₹$€£]?\s*\(?[\d,]+(?:\.\d+)?\s*(?:M|million|B|billion|K|thousand|Cr|L)?\)?)*)$",
            re.MULTILINE | re.IGNORECASE
        )

        for line in lines:
            line = line.strip()
            if not line:
                continue

            cleaned_line = (
                line.replace(chr(8216), "'")
                .replace(chr(8217), "'")
                .replace(chr(8220), '"')
                .replace(chr(8221), '"')
            )
            cleaned_line = re.sub(r"(\d+),\s+(\d{3})", r"\1,\2", cleaned_line)
            cleaned_line = re.sub(r"([$€£₹])\s+(\(?[\d,]+(?:\.\d+)?)", r"\1\2", cleaned_line)

            parts = re.split(r"\s{2,}|\t", cleaned_line)
            if len(parts) >= 2 and any(re.search(r"\d", p) for p in parts[1:]):
                label = parts[0].strip()
                vals = [p.strip() for p in parts[1:]]
                if should_reverse_cols:
                    vals = list(reversed(vals))
                rows.append([label] + vals)
                continue

            match = line_pattern.match(cleaned_line)
            if match:
                particular = match.group(1).strip()
                vals_raw = match.group(2).strip()
                val_tokens = [
                    v.strip()
                    for v in re.split(r"\s+", vals_raw)
                    if v.strip() and v.strip() not in ("$", "€", "£", "₹")
                ]
                if val_tokens:
                    if should_reverse_cols:
                        val_tokens = list(reversed(val_tokens))
                    rows.append([particular] + val_tokens)

        if detected_headers and rows:
            headers = detected_headers
        else:
            headers = ["Particulars"] + [
                f"Value_{i+1}" for i in range(max((len(r) - 1 for r in rows), default=1))
            ]

        return RawTable(headers=headers, rows=rows, page_number=1)
