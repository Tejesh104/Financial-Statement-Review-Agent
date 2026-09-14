import pandas as pd
import pymupdf as fitz
from pathlib import Path

def create_sample_files():
    # 1. Create sample_financial.csv
    csv_content = (
        'ABC Limited - Statement of Accounts FY 2024-25 (INR),Amount\n'
        'Total Revenue,"12,500,000"\n'
        'Total Assets,"45,000,000"\n'
        'Total Liabilities,"18,000,000"\n'
        'Shareholders\' Equity,"27,000,000"\n'
    )
    Path("sample_financial.csv").write_text(csv_content, encoding="utf-8")
    print("Created sample_financial.csv")

    # 2. Create sample_financial.xlsx
    excel_df = pd.DataFrame({
        "Particulars": [
            "Revenue",
            "Total Assets",
            "Total Liabilities",
            "Owners' Equity"
        ],
        "FY 2024-25 (₹)": [
            "12,500,000",
            "45,000,000",
            "18,000,000",
            "27,000,000"
        ]
    })
    # Add company banner
    with pd.ExcelWriter("sample_financial.xlsx", engine="openpyxl") as writer:
        excel_df.to_excel(writer, sheet_name="Balance_Sheet", index=False)
    
    # Prepend company title line to excel workbook using openpyxl
    import openpyxl
    wb = openpyxl.load_workbook("sample_financial.xlsx")
    ws = wb["Balance_Sheet"]
    ws.insert_rows(1)
    ws.cell(row=1, column=1, value="ABC Limited")
    ws.cell(row=1, column=2, value="INR")
    wb.save("sample_financial.xlsx")
    print("Created sample_financial.xlsx")

    # 3. Create sample_financial.pdf
    doc = fitz.open()
    page = doc.new_page()
    pdf_text = (
        "ABC Limited\n"
        "Financial Statement for FY 2024-25\n"
        "Currency: INR (₹)\n\n"
        "Particulars                     FY 2024-25\n"
        "Turnover                        12,500,000\n"
        "Total Assets                    45,000,000\n"
        "Total Liabilities               18,000,000\n"
        "Shareholders Equity             27,000,000\n"
    )
    page.insert_text((50, 72), pdf_text, fontsize=12)
    doc.save("sample_financial.pdf")
    doc.close()
    print("Created sample_financial.pdf")

if __name__ == "__main__":
    create_sample_files()
